"""
Reads bulk score files (.csv, .xlsx) from config.SCORE_IMPORT_DIRECTORY and
parses them into ScoreImportRow entries ready for
AssessmentRepository.bulk_upsert().

Unlike RosterImportService (which reads plain, headerless roster files and
relies on column position - name(s) then a trailing Sid), score files are
expected to come from an export: either this app's own "Export to CSV /
Export to Excel" output, another TA's export of the same shape, or an
autograder's report. Those always have a header row, so columns here are
matched BY NAME (case-insensitive) instead of by position. This is more
robust to files that have extra columns this app doesn't care about (e.g.
a "Calculated" or "Last Modified" column from one of our own exports) -
anything not recognized is simply ignored.

Required column: Sid (or "student id").
Recognized optional columns:
  Score   ("score")            - the grade itself.
  Comment ("comment")          - free-text comment for that Sid.
  Grader  ("grader")           - a TA name to attribute the imported score
                                  to, e.g. from one of our own exports.
                                  If absent, the caller (MainWindow) falls
                                  back to whichever TA is running the
                                  import. Resolving the name to a TaId is
                                  the caller's job (via TARepository), not
                                  this service's - this module only reads
                                  and parses the file.
Every other column (SName, Calculated, Last Modified, ...) is ignored.
"""
import csv
import os
from dataclasses import dataclass
from typing import List, Optional

from openpyxl import load_workbook

from config import SCORE_IMPORT_DIRECTORY
from logging_setup import get_logger

logger = get_logger("services.score_import")

SUPPORTED_EXTENSIONS = (".csv", ".xlsx")

_SID_HEADERS = {"sid", "studentid", "student id", "id"}
_SCORE_HEADERS = {"score"}
_COMMENT_HEADERS = {"comment"}
_GRADER_HEADERS = {"grader", "graderid", "grader id"}


@dataclass
class ScoreImportRow:
    sid: int
    score: Optional[float]
    comment: str = "-"
    grader_name: Optional[str] = None


class ScoreImportError(Exception):
    pass


class ScoreImportService:
    def __init__(self, import_directory: str = SCORE_IMPORT_DIRECTORY):
        self.import_directory = import_directory

    def list_score_files(self) -> List[str]:
        """Filenames available to pick from at bulk-import time."""
        if not os.path.isdir(self.import_directory):
            return []
        return sorted(
            f for f in os.listdir(self.import_directory)
            if f.lower().endswith(SUPPORTED_EXTENSIONS)
        )

    def parse_score_file(self, filename: str) -> List[ScoreImportRow]:
        full_path = os.path.join(self.import_directory, filename)
        if not os.path.isfile(full_path):
            raise ScoreImportError(f"Score file not found: {full_path}")

        ext = os.path.splitext(filename)[1].lower()
        if ext == ".csv":
            rows = self._parse_csv(full_path, filename)
        elif ext == ".xlsx":
            rows = self._parse_excel(full_path, filename)
        else:
            raise ScoreImportError(
                f"Unsupported score file type '{ext}'. "
                f"Supported types: {', '.join(SUPPORTED_EXTENSIONS)}"
            )

        if not rows:
            logger.warning("No valid score rows found in %s", filename)
        return rows

    # ---- format-specific parsers ----
    def _parse_csv(self, full_path: str, filename: str) -> List[ScoreImportRow]:
        with open(full_path, "r", encoding="utf-8", newline="") as f:
            reader = csv.reader(f)
            all_rows = list(reader)
        return self._parse_tabular(all_rows, filename)

    def _parse_excel(self, full_path: str, filename: str) -> List[ScoreImportRow]:
        try:
            workbook = load_workbook(full_path, read_only=True, data_only=True)
        except Exception as e:
            raise ScoreImportError(f"Could not open Excel file '{filename}': {e}") from e

        try:
            sheet = workbook.active
            all_rows = [
                ["" if v is None else str(v).strip() for v in row]
                for row in sheet.iter_rows(values_only=True)
            ]
        finally:
            workbook.close()
        return self._parse_tabular(all_rows, filename)

    # ---- shared helpers ----
    def _parse_tabular(self, all_rows, filename: str) -> List[ScoreImportRow]:
        if not all_rows:
            return []

        header = [str(c).strip().lower() for c in all_rows[0]]
        sid_col = self._find_column(header, _SID_HEADERS)
        if sid_col is None:
            raise ScoreImportError(
                f"'{filename}' has no recognizable Sid column. "
                f"Expected a header named one of: {', '.join(sorted(_SID_HEADERS))}."
            )
        score_col = self._find_column(header, _SCORE_HEADERS)
        comment_col = self._find_column(header, _COMMENT_HEADERS)
        grader_col = self._find_column(header, _GRADER_HEADERS)

        rows: List[ScoreImportRow] = []
        for row_number, raw_row in enumerate(all_rows[1:], start=2):
            sid_str = self._cell(raw_row, sid_col)
            sid = self._try_parse_sid(sid_str)
            if sid is None:
                logger.warning(
                    "Skipping row %d in %s: Sid value %r is not valid",
                    row_number, filename, sid_str,
                )
                continue

            score = self._try_parse_score(self._cell(raw_row, score_col)) if score_col is not None else None
            comment = self._cell(raw_row, comment_col) if comment_col is not None else ""
            grader_name = self._cell(raw_row, grader_col) if grader_col is not None else None
            grader_name = grader_name or None  # normalize "" -> None

            rows.append(ScoreImportRow(sid=sid, score=score, comment=comment or "-", grader_name=grader_name))
        return rows

    @staticmethod
    def _find_column(header: List[str], candidates: set) -> Optional[int]:
        for index, name in enumerate(header):
            if name in candidates:
                return index
        return None

    @staticmethod
    def _cell(row, index: Optional[int]) -> str:
        if index is None or index >= len(row):
            return ""
        return str(row[index]).strip()

    @staticmethod
    def _try_parse_sid(value: str):
        try:
            return int(value)
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _try_parse_score(value: str):
        if value == "":
            return None
        try:
            return float(value) if "." in value else int(value)
        except (TypeError, ValueError):
            return None